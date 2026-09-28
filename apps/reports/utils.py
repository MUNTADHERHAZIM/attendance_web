import os
import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
from django.contrib.auth import get_user_model
from django.utils import timezone
from apps.accounts.models import StudentProfile

User = get_user_model()

def import_students_from_excel(file_path, institution, default_shift="MORNING"):
    """
    Parses an Excel or CSV file with ULTRA-FLEXIBLE structure:
    - Even a SINGLE COLUMN with student names is 100% valid!
    - Full name is automatically split into first_name and last_name.
    - Missing student_id, username, email are automatically generated.
    - All other columns (phone, email, shift) are completely optional.
    """
    import random
    import time

    errors = []
    valid_students = []
    
    try:
        if file_path.endswith(".csv"):
            df = pd.read_csv(file_path)
        else:
            df = pd.read_excel(file_path)
    except Exception as e:
        errors.append(f"تعذر قراءة الملف: {str(e)}")
        return [], errors

    if df.empty:
        errors.append("الملف فارغ ولا يحتوي على أي بيانات.")
        return [], errors

    # Clean columns
    df.columns = [str(c).strip() for c in df.columns]

    # Map columns intelligently
    col_name = None
    col_id = None
    col_phone = None
    col_email = None
    col_shift = None
    col_username = None

    for col in df.columns:
        c_low = col.lower()
        if not col_name and any(k in c_low for k in ["اسم", "name", "طالب", "student", "full"]):
            col_name = col
        elif not col_id and any(k in c_low for k in ["رقم", "جامعي", "student_id", "id", "code"]):
            col_id = col
        elif not col_phone and any(k in c_low for k in ["هاتف", "موبايل", "phone", "mobile"]):
            col_phone = col
        elif not col_email and any(k in c_low for k in ["بريد", "email", "mail"]):
            col_email = col
        elif not col_shift and any(k in c_low for k in ["دراسة", "shift"]):
            col_shift = col
        elif not col_username and any(k in c_low for k in ["مستخدم", "username"]):
            col_username = col

    # Fallback: if no name column matched, use the first column!
    if not col_name and len(df.columns) > 0:
        col_name = df.columns[0]

    current_year = timezone.now().year
    base_counter = StudentProfile.objects.count() + 1001

    for index, row in df.iterrows():
        row_num = index + 2
        
        # 1. Full Name
        raw_name = str(row[col_name]).strip() if col_name and pd.notna(row.get(col_name)) else ""
        if not raw_name or raw_name.lower() in ["nan", "none", "null"]:
            continue

        parts = raw_name.split()
        first_name = parts[0]
        last_name = " ".join(parts[1:]) if len(parts) > 1 else parts[0]

        # 2. Student ID
        raw_id = str(row[col_id]).strip() if col_id and pd.notna(row.get(col_id)) else ""
        if not raw_id or raw_id.lower() in ["nan", "none", "null"]:
            raw_id = f"S{current_year}{(base_counter + index):04d}"

        # 3. Username
        raw_username = str(row[col_username]).strip() if col_username and pd.notna(row.get(col_username)) else ""
        if not raw_username or raw_username.lower() in ["nan", "none", "null"]:
            raw_username = f"std_{raw_id.lower().replace(' ', '_')}"

        # 4. Email
        raw_email = str(row[col_email]).strip() if col_email and pd.notna(row.get(col_email)) else ""
        if not raw_email or raw_email.lower() in ["nan", "none", "null"]:
            raw_email = f"{raw_username}@student.uobaghdad.edu.iq"

        # 5. Phone
        raw_phone = str(row[col_phone]).strip() if col_phone and pd.notna(row.get(col_phone)) else ""
        if raw_phone.lower() in ["nan", "none", "null"]:
            raw_phone = ""

        # 6. Study Shift
        shift_val = default_shift
        if col_shift and pd.notna(row.get(col_shift)):
            s_str = str(row[col_shift]).strip()
            if "مسائ" in s_str or "even" in s_str.lower():
                shift_val = "EVENING"
            elif "صباح" in s_str or "morn" in s_str.lower():
                shift_val = "MORNING"

        valid_students.append({
            "full_name": raw_name,
            "first_name": first_name,
            "last_name": last_name,
            "student_id": raw_id,
            "username": raw_username,
            "email": raw_email,
            "phone": raw_phone,
            "study_shift": shift_val
        })

    return valid_students, errors


def parse_students_from_text(raw_text, default_shift="MORNING"):
    """
    Parses a plain-text list of student names (one per line).
    """
    valid_students = []
    current_year = timezone.now().year
    base_counter = StudentProfile.objects.count() + 1001

    lines = [line.strip() for line in raw_text.splitlines() if line.strip()]
    for index, line in enumerate(lines):
        # Ignore common table header lines
        if any(h in line for h in ["اسم الطالب", "الاسم الكامل", "الاسم الثلاثي", "ت", "الرقم"]):
            continue

        parts = line.split()
        if not parts:
            continue

        first_name = parts[0]
        last_name = " ".join(parts[1:]) if len(parts) > 1 else parts[0]
        student_id = f"S{current_year}{(base_counter + index):04d}"
        username = f"std_{student_id.lower()}"
        email = f"{username}@student.uobaghdad.edu.iq"

        valid_students.append({
            "full_name": line,
            "first_name": first_name,
            "last_name": last_name,
            "student_id": student_id,
            "username": username,
            "email": email,
            "phone": "",
            "study_shift": default_shift
        })

    return valid_students


def export_attendance_excel(attendance_session, records):
    """
    Creates an Excel spreadsheet (Workbook object) for the class session attendance.
    """
    wb = Workbook()
    ws = wb.active
    ws.title = "تقرير الحضور"
    
    # Configure Sheet to be RTL (Right to Left) for Arabic layout
    ws.sheet_view.rightToLeft = True
    
    # Style definitions
    font_title = Font(name="Segoe UI", size=16, bold=True, color="0F172A")
    font_subtitle = Font(name="Segoe UI", size=11, bold=True, color="475569")
    font_header = Font(name="Segoe UI", size=11, bold=True, color="FFFFFF")
    font_body = Font(name="Segoe UI", size=10)
    
    fill_header = PatternFill(start_color="1E293B", end_color="1E293B", fill_type="solid")
    fill_present = PatternFill(start_color="DCFCE7", end_color="DCFCE7", fill_type="solid")  # light green
    fill_absent = PatternFill(start_color="FEE2E2", end_color="FEE2E2", fill_type="solid")   # light red
    fill_late = PatternFill(start_color="FEF9C3", end_color="FEF9C3", fill_type="solid")     # light yellow
    fill_excused = PatternFill(start_color="F1F5F9", end_color="F1F5F9", fill_type="solid")  # light grey

    thin_border = Border(
        left=Side(style="thin", color="CBD5E1"),
        right=Side(style="thin", color="CBD5E1"),
        top=Side(style="thin", color="CBD5E1"),
        bottom=Side(style="thin", color="CBD5E1")
    )
    
    # Add metadata headers
    course = attendance_session.session.course
    section = attendance_session.session.class_section
    
    ws.append(["نظام الحضور الذكي - تقرير الحضور والغياب"])
    ws.cell(row=1, column=1).font = font_title
    
    ws.append([f"المادة: {course.name} ({course.code})", f"الشعبة: {section.level} - {section.name}"])
    ws.append([f"المعلم: {attendance_session.session.teacher.user.get_full_name()}", f"التاريخ: {attendance_session.date}"])
    
    for row in [2, 3]:
        for col in [1, 2]:
            ws.cell(row=row, column=col).font = font_subtitle
            
    ws.append([])  # Spacer row
    
    # Table headers
    headers = ["الرقم الجامعي", "اسم الطالب", "حالة الحضور", "طريقة التحضير", "الوقت", "عنوان الـ IP", "ملاحظات"]
    ws.append(headers)
    
    header_row_idx = 5
    for col_idx in range(1, len(headers) + 1):
        cell = ws.cell(row=header_row_idx, column=col_idx)
        cell.font = font_header
        cell.fill = fill_header
        cell.alignment = Alignment(horizontal="center", vertical="center")
        cell.border = thin_border
        
    # Add records
    for rec in records:
        student = rec.student
        status_text = rec.get_status_display()
        method_text = rec.get_method_display()
        timestamp_text = timezone.localtime(rec.timestamp).strftime("%H:%M:%S") if rec.timestamp else "-"
        ip_addr = rec.ip_address or "-"
        notes = rec.notes or ""
        
        row_data = [
            student.student_id,
            student.user.get_full_name(),
            status_text,
            method_text,
            timestamp_text,
            ip_addr,
            notes
        ]
        ws.append(row_data)
        
        # Color-code status cell
        curr_row = ws.max_row
        status_cell = ws.cell(row=curr_row, column=3)
        if rec.status == "PRESENT":
            status_cell.fill = fill_present
        elif rec.status == "ABSENT":
            status_cell.fill = fill_absent
        elif rec.status == "LATE":
            status_cell.fill = fill_late
        else:
            status_cell.fill = fill_excused
            
        # Apply standard borders and alignment
        for col_idx in range(1, len(headers) + 1):
            cell = ws.cell(row=curr_row, column=col_idx)
            cell.font = font_body
            cell.border = thin_border
            cell.alignment = Alignment(horizontal="right" if col_idx == 2 else "center")
            
    # Auto-fit columns
    for col in ws.columns:
        max_len = max(len(str(cell.value or '')) for cell in col)
        col_letter = col[0].column_letter
        ws.column_dimensions[col_letter].width = max(max_len + 3, 12)
        
    return wb
